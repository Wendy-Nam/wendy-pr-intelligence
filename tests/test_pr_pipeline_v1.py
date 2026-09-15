from prmonitor.pipelines.pr import to_pr_records
from prmonitor.renderers.exports import export_csv
def test_pr_records_and_csv_use_same_records():
 records=to_pr_records([{'id':'a','title':'Our Robot update','summary':'news','source_name':'S'}],company_aliases=['Our Robot'])
 assert records[0]['tone']=='neutral' and export_csv(records).count('\n')==2
def test_csv_formula_values_are_neutralized():
 assert "'=1" in export_csv([{'id':'a','title':'=1','summary':'','tone':'','source_name':'','source_url':''}])
def test_xlsx_uses_same_record_rows(tmp_path):
 from prmonitor.renderers.exports import export_xlsx
 path=export_xlsx([{'id':'a','title':'T','summary':'S','tone':'neutral','source_name':'N','source_url':'https://x'}],tmp_path/'out.xlsx')
 from openpyxl import load_workbook
 assert load_workbook(path).active.max_row == 2
def test_pr_html_rejects_unsafe_link_and_escapes_title():
 from prmonitor.renderers.pr import render_pr_html
 html=render_pr_html([{'title':'<x>','summary':'s','source_url':'javascript:bad'}])
 assert '&lt;x&gt;' in html and 'javascript:' not in html
